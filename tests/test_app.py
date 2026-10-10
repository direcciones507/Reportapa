import os, io, uuid
from datetime import datetime, timezone
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from app import main as m

@pytest.fixture
def data():
    return dict(name='Familia Prueba',phone='+507 6123-4567',province='Veraguas',district='Santiago',township='Santiago',community='Sector prueba',maps='https://maps.app.goo.gl/abc123',damage='Daño de prueba',needs=['Agua potable'])

def test_validation(data):
    assert m.validate(data)['phone']=='61234567'
    for link in ['https://google.com.evil.test/maps','http://maps.app.goo.gl/a','https://google.com/search','https://user@google.com/maps','https://goo.gl/anything']:
        with pytest.raises(HTTPException): m.validate({**data,'maps':link})
    with pytest.raises(HTTPException): m.validate({**data,'damage':''})

def test_images():
    out=io.BytesIO();Image.new('RGB',(20,20)).save(out,'PNG');result=m.clean_photo(out.getvalue())
    assert Image.open(io.BytesIO(result)).format=='JPEG'
    with pytest.raises(HTTPException): m.clean_photo(b'fake image')
    with pytest.raises(HTTPException): m.clean_photo(b'x'*(5*1024*1024+1))

def test_excel(data):
    data=m.validate(data);data['name']='=HYPERLINK("evil")';data['damage']='secret damage';data['needs_detail']='private detail'
    rows=[dict(id='RPA-00000001',created_at=datetime.now(timezone.utc),data=data,status='pendiente')]
    wb=load_workbook(io.BytesIO(m.workbook(rows,'todos')))
    assert wb['Reportes']['C2'].data_type=='s'
    assert wb['Información']['B3'].value==1
    anonymous=load_workbook(io.BytesIO(m.workbook(rows,'todos',True)))
    text=str(list(anonymous['Reportes'].values))
    for private in ['HYPERLINK','61234567','abc123','secret damage','private detail','Sector prueba']: assert private not in text

def test_closed_and_auth(monkeypatch):
    monkeypatch.setenv('SESSION_SECRET','s'*40)
    monkeypatch.delenv('PUBLIC_REPORTS_ENABLED',raising=False)
    client=TestClient(m.app)
    assert client.post('/api/reports').status_code==503
    for path in ['/api/admin/reports','/api/admin/export/master','/api/admin/export/area','/api/admin/export/anonymous','/api/admin/export/summary','/api/admin/reports/RPA-1/photos/file.jpg']:
        assert client.get(path).status_code==401
    assert client.get('/').status_code==200
    assert client.get('/').headers['cache-control']=='no-store'

@pytest.fixture
def live(monkeypatch,tmp_path):
    if not os.getenv('DATABASE_URL'): pytest.skip('PostgreSQL required')
    monkeypatch.setenv('SESSION_SECRET','test-secret-'*4)
    monkeypatch.setenv('ADMIN_PASSWORD','test-password-123456')
    monkeypatch.setenv('PHOTO_DIR',str(tmp_path))
    monkeypatch.setenv('COOKIE_SECURE','false')
    monkeypatch.setenv('PUBLIC_REPORTS_ENABLED','true')
    with TestClient(m.app) as client:
        with m.db() as c: c.execute('TRUNCATE reports,request_limits')
        yield client

def submit(client,data,key=None,extra=None,photos=None):
    fields={**data,'consent':'yes','idempotency_key':key or str(uuid.uuid4()),**(extra or {})}
    fields.pop('needs');fields['needs']='Agua potable'
    return client.post('/api/reports',data=fields,files=photos or {'photos':('',b'','application/octet-stream')})

def test_postgres_flow(live,data):
    key=str(uuid.uuid4());r=submit(live,data,key);assert r.status_code==200,r.text;identifier=r.json()['id'];assert identifier.startswith('RPA-')
    assert submit(live,data,key).json()['id']==identifier
    assert submit(live,{**data,'damage':'changed'},key).status_code==409
    dup=submit(live,data);assert dup.status_code==409 and 'id' not in dup.json()
    assert submit(live,data,extra={'new_damage':'yes'}).status_code==200
    login=live.post('/api/admin/login',json={'password':'test-password-123456'});assert login.status_code==200
    csrf={'x-csrf-token':login.json()['csrf']}
    assert live.get('/api/admin/reports').json()['total']==2
    assert live.get('/api/admin/reports?province=Coclé').json()['total']==0
    assert live.patch('/api/admin/reports/'+identifier,json={'status':'visitado'}).status_code==403
    assert live.patch('/api/admin/reports/'+identifier,json={'status':'visitado'},headers=csrf).status_code==200
    assert live.get('/api/admin/reports?status=visitado&need=Agua%20potable').json()['total']==1
    r=live.get('/api/admin/export/master');assert r.status_code==200
    wb=load_workbook(io.BytesIO(r.content));assert wb['Reportes'].max_row==3
    assert live.get('/api/admin/export/summary').content.startswith(b'%PDF')
    assert live.post('/api/admin/logout',headers=csrf).status_code==200
    assert live.get('/api/admin/reports').status_code==401

def test_private_photos(live,data):
    out=io.BytesIO();Image.new('RGB',(20,20)).save(out,'PNG')
    r=submit(live,data,photos={'photos':('test.png',out.getvalue(),'image/png')});assert r.status_code==200,r.text
    live.post('/api/admin/login',json={'password':'test-password-123456'})
    report=live.get('/api/admin/reports').json()['reports'][0];url='/api/admin/reports/'+report['id']+'/photos/'+report['photos'][0]
    assert live.get(url).status_code==200
    assert TestClient(m.app).get(url).status_code==401
    assert live.get('/static/'+report['photos'][0]).status_code==404

def test_rate_limit(live):
    for _ in range(10): assert live.post('/api/admin/login',json={'password':'wrong'}).status_code==401
    assert live.post('/api/admin/login',json={'password':'wrong'}).status_code==429
