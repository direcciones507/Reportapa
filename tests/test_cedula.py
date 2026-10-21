"""Private responsible ID: validation, storage, export privacy and legacy compatibility."""
import io
import json
from datetime import datetime, timezone
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader
from app import main as m
from app.report_data import matching_keys
from test_app import data, live, submit

@pytest.mark.parametrize('value', ['8-123-456', 'PE-12-3456', 'E-123-456', '8AV-1-123', 'N-1-123'])
def test_valid_cedula(data, value):
    assert m.validate({**data, 'cedula': ' '+value.lower()+' '})['cedula'] == value

@pytest.mark.parametrize('value', [123, True, [], {}, '<script>', '=1+1', '8/123/456', '8-1-1\nBAD', 'A'*21, '8-1234567-1'])
def test_invalid_cedula(data, value):
    with pytest.raises(HTTPException) as error:
        m.validate({**data, 'cedula': value})
    assert error.value.status_code == 422

@pytest.mark.parametrize('value', [None, '', '  '])
def test_optional_cedula_does_not_change_legacy_payload(data, value):
    assert m.validate({**data, 'cedula': value}) == m.validate(data)


def test_cedula_does_not_change_dedup(data):
    original=m.validate(data)
    updated=m.validate({**data, 'cedula':'8-123-456'})
    assert matching_keys(original, lambda s:s) == matching_keys(updated, lambda s:s)


def test_cedula_private_exports_and_legacy(data, tmp_path):
    report=dict(id='RPA-TEST', created_at=datetime.now(timezone.utc),
                data=m.validate({**data, 'cedula':'8-123-456'}), status='pendiente')
    for personal, anonymous in [(True,False), (False,False), (True,True), (False,True)]:
        wb=load_workbook(io.BytesIO(m.workbook([report], 'test', anonymous, personal)))
        all_cells=str([list(sheet.values) for sheet in wb])
        assert ('8-123-456' in all_cells) == (personal and not anonymous)
        assert ('Número de cédula' in all_cells) == (personal and not anonymous)
    pdf=PdfReader(io.BytesIO(m.individual_pdf(report,tmp_path,m.HELP_STATES)))
    assert len(pdf.pages)==1 and '8-123-456' in pdf.pages[0].extract_text()
    del report['data']['cedula']
    assert 'Sin informar' in PdfReader(io.BytesIO(m.individual_pdf(report,tmp_path,m.HELP_STATES))).pages[0].extract_text()
    wb=load_workbook(io.BytesIO(m.workbook([report],'test')))
    values=list(wb['Detalle de reportes'].values)
    assert values[1][values[0].index('Número de cédula')] is None


def test_cedula_postgres_roundtrip_and_migration(live, data):
    # A historical row remains byte-for-byte identical after reapplying migration.
    old=submit(live,data).json()['id']
    with m.db() as c:
        before=c.execute('SELECT to_jsonb(r) AS row FROM reports r WHERE id=%s',(old,)).fetchone()['row']
        migration=(m.ROOT.parent/'migrations/008_responsible_cedula.sql').read_text()
        c.execute(migration); c.execute(migration)
        after=c.execute('SELECT to_jsonb(r) AS row FROM reports r WHERE id=%s',(old,)).fetchone()['row']
        assert before==after and 'cedula' not in after['data']
        assert c.execute("SELECT convalidated FROM pg_constraint WHERE conname='reports_cedula_format'").fetchone()['convalidated']
    result=submit(live,{**data,'phone':'69998877','family_name':'Otra familia','cedula':'8-123-456'})
    assert result.status_code==200,result.text
    identifier=result.json()['id']
    assert 'cedula' not in result.text and '8-123-456' not in result.text
    assert live.get('/api/admin/reports').status_code==401
    assert live.get('/api/admin/reports/'+identifier+'/pdf').status_code==401
    login=live.post('/api/admin/login',json={'username':'Ayudapanama','password':'test-password-123456'})
    assert login.status_code==200
    reports=live.get('/api/admin/reports').json()['reports']
    assert next(r for r in reports if r['id']==identifier)['data']['cedula']=='8-123-456'
    with m.db() as c:
        assert c.execute('SELECT data FROM reports WHERE id=%s',(identifier,)).fetchone()['data']['cedula']=='8-123-456'
    pdf=PdfReader(io.BytesIO(live.get('/api/admin/reports/'+identifier+'/pdf').content))
    assert '8-123-456' in pdf.pages[0].extract_text()
    for kind in ('master','area','anonymous'):
        response=live.get('/api/admin/export/'+kind+'?include_personal=true')
        assert response.status_code==200
        wb=load_workbook(io.BytesIO(response.content))
        assert ('8-123-456' in str([list(s.values) for s in wb])) == (kind!='anonymous')
    assert submit(live,{**data,'cedula':'9-111-222'}).status_code==409
