import io, uuid, json, shutil
from pathlib import Path
from datetime import datetime,timezone
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from app import main as m
from app.report_data import household,coordinates,matching_keys
from app.photo_maintenance import cleanup_photos
from test_app import live, data, submit

def login(client):
    r=client.post('/api/admin/login',json={'username':'Ayudapanama','password':'test-password-123456'})
    assert r.status_code==200
    return {'x-csrf-token':r.json()['csrf']}

def test_household_counts(data):
    value=m.validate(data);assert value['total_people']==4 and value['age_0_5'] is None
    assert household({**data,'age_0_5':1,'age_6_17':0,'age_18_59':2,'age_60_plus':1,'people_disability':1})['age_6_17']==0
    # Partial information is permitted, but never overcounts the whole family.
    assert household({**data,'age_0_5':1})['age_18_59'] is None
    invalid=[{'total_people':0},{'total_people':1001},{'total_people':True},{'total_people':'1.5'},{'family_name':''},{'age_0_5':-1},{'age_0_5':5},{'age_0_5':1,'age_6_17':1,'age_18_59':0,'age_60_plus':0},{'people_disability':5}]
    for extra in invalid:
        with pytest.raises(HTTPException):m.validate({**data,**extra})

def test_coordinates_only_when_explicit():
    for link in ['https://www.google.com/maps?q=8.12345,-80.98765','https://www.google.com/maps/search/?api=1&query=8.12345%2C-80.98765','https://www.google.com/maps/place/Test/data=!3d8.12345!4d-80.98765']:
        assert coordinates(link)=={'latitude':8.12345,'longitude':-80.98765}
    for link in ['https://maps.app.goo.gl/abc','https://www.google.com/maps/@8.1,-80.9,15z','https://www.google.com/maps?ll=8.1,-80.9','https://www.google.com/maps?q=91,-80','https://www.google.com/maps?q=Vivienda']:
        assert coordinates(link) is None

def test_household_exports_and_pdf(data,tmp_path):
    d=m.validate({**data,'age_0_5':1,'age_6_17':0,'age_18_59':2,'age_60_plus':1,'people_disability':1,'maps':'https://www.google.com/maps?q=8.1,-80.9'})
    r=dict(id='RPA-00000001',created_at=datetime.now(timezone.utc),data=d,status='pendiente',photos=[])
    wb=load_workbook(io.BytesIO(m.workbook([r],'todos')));values=list(wb['Reportes'].values);row=dict(zip(*values))
    assert row['Nombre de familia']=='Familia Prueba' and row['Total de personas afectadas']==4
    assert row['0 a 5 años']==1 and row['Personas con discapacidad']==1 and row['Latitud']==8.1
    anon=load_workbook(io.BytesIO(m.workbook([r],'todos',True)));anonymous=str(list(anon['Reportes'].values))
    for secret in ['Familia Prueba','Responsable Prueba','61234567','8.1','-80.9']:assert secret not in anonymous
    text=PdfReader(io.BytesIO(m.individual_pdf(r,tmp_path,m.HELP_STATES))).pages[0].extract_text()
    assert 'Familia Prueba' in text and 'Total de personas afectadas' in text and '0 a 5 años: 1' in text
    # Legacy records remain readable without fabricated household data.
    for field in ['family_name','total_people',*m.COUNT_FIELDS]:d.pop(field,None)
    legacy=PdfReader(io.BytesIO(m.individual_pdf(r,tmp_path,m.HELP_STATES)))
    assert len(legacy.pages)==1 and 'Sin informar' in legacy.pages[0].extract_text()

def test_stronger_duplicates(live,data):
    maps='https://www.google.com/maps?q=8.12345,-80.98765'
    first={**data,'maps':maps};key=str(uuid.uuid4());assert submit(live,first,key).status_code==200
    # Same submission and key succeeds idempotently; new key cannot repeat content.
    assert submit(live,first,key).status_code==200
    identical=submit(live,first,extra={'new_damage':'yes'});assert identical.status_code==409 and 'id' not in identical.json()
    changed={**first,'maps':'https://www.google.com/maps/search/?query=8.12345,-80.98765','damage':'Daño nuevo'}
    probable=submit(live,changed);assert probable.status_code==409 and probable.json()['possible_duplicate']
    for secret in ['Familia Prueba','61234567','RPA-']:assert secret not in probable.text
    assert submit(live,changed,extra={'new_damage':'yes'}).status_code==200
    other={**changed,'family_name':'Otra familia','damage':'Pared distinta'}
    assert submit(live,other).status_code==409
    assert submit(live,other,extra={'new_damage':'yes'}).status_code==200
    # Same household + sector also warns after phone/Maps link change.
    switched={**first,'phone':'69876543','maps':'https://maps.app.goo.gl/new','damage':'Otra descripción'}
    assert submit(live,switched).json()['possible_duplicate']

def test_archive_and_permanent_delete(live,data):
    out=io.BytesIO();Image.new('RGB',(20,20)).save(out,'PNG')
    identifier=submit(live,data,photos={'photos':('photo.png',out.getvalue(),'image/png')}).json()['id'];url='/api/admin/reports/'+identifier
    assert live.request('DELETE',url,json={'confirm':'ELIMINAR '+identifier}).status_code==401
    csrf=login(live);report=live.get('/api/admin/reports').json()['reports'][0]
    path=m.photo_dir()/report['photos'][0]
    assert live.patch(url,json={'archived':'yes'},headers=csrf).status_code==422
    assert live.patch(url,json={'archived':True},headers=csrf).status_code==200
    assert live.get('/api/admin/reports').json()['total']==0
    archived=live.get('/api/admin/reports?archive=archived').json()['reports'][0]
    assert archived['data']==report['data'] and archived['photos']==report['photos'] and path.exists()
    assert live.get(url+'/pdf').status_code==200
    assert live.get(url+'/photos/'+path.name).status_code==200
    master=load_workbook(io.BytesIO(live.get('/api/admin/export/master').content));assert master['Reportes'].max_row==2
    area=load_workbook(io.BytesIO(live.get('/api/admin/export/area?archive=archived').content));assert area['Reportes'].max_row==2
    assert live.patch(url,json={'archived':False},headers=csrf).status_code==200
    assert live.get('/api/admin/reports').json()['total']==1
    assert live.request('DELETE',url,json={'confirm':'ELIMINAR '+identifier}).status_code==403
    assert live.request('DELETE',url,json={'confirm':'ELIMINAR RPA-OTRO'},headers=csrf).status_code==422
    result=live.request('DELETE',url,json={'confirm':'ELIMINAR '+identifier},headers=csrf)
    assert result.status_code==200 and result.json()['photos_pending']==0
    assert not path.exists() and live.get('/api/admin/reports?archive=all').json()['total']==0
    assert live.get(url+'/pdf').status_code==404 and live.get(url+'/photos/'+path.name).status_code==404
    with m.db() as c:
        assert c.execute('SELECT count(*) AS n FROM reports').fetchone()['n']==0
        assert c.execute('SELECT count(*) AS n FROM photo_cleanup').fetchone()['n']==0
    ledger=live.get('/api/admin/export/erasures');assert ledger.status_code==200
    assert ledger.json()[0]['id']==identifier and 'phone' not in ledger.text and 'Familia' not in ledger.text
    assert TestClient(m.app).get('/api/admin/export/erasures').status_code==401

def test_physical_erasure_failure_retry(live,data,monkeypatch):
    out=io.BytesIO();Image.new('RGB',(20,20)).save(out,'PNG')
    id=submit(live,data,photos={'photos':('photo.png',out.getvalue(),'image/png')}).json()['id'];csrf=login(live)
    report=live.get('/api/admin/reports').json()['reports'][0];name=report['photos'][0];path=m.photo_dir()/name
    original=Path.unlink
    def denied(self,*args,**kwargs):
        if self.name==name:raise PermissionError('simulated volume error')
        return original(self,*args,**kwargs)
    with monkeypatch.context() as patch:
        patch.setattr(Path,'unlink',denied)
        result=live.request('DELETE','/api/admin/reports/'+id,json={'confirm':'ELIMINAR '+id},headers=csrf)
        assert result.status_code==202 and result.json()['photos_pending']==1
        assert path.exists() and live.get('/api/admin/reports').json()['total']==0
        assert live.get('/api/admin/reports/'+id+'/photos/'+name).status_code==404
    cleanup_photos(m.db,m.photo_dir())
    assert not path.exists()
    with m.db() as c:assert c.execute('SELECT count(*) AS n FROM photo_cleanup').fetchone()['n']==0

def test_volume_restart_and_backup_recovery(live,data,tmp_path):
    out=io.BytesIO();Image.new('RGB',(20,20),(12,120,200)).save(out,'PNG')
    id=submit(live,data,photos={'photos':('photo.png',out.getvalue(),'image/png')}).json()['id'];csrf=login(live)
    with m.db() as c:record=c.execute('SELECT * FROM reports WHERE id=%s',(id,)).fetchone()
    path=m.photo_dir()/record['photos'][0];backup=path.read_bytes()
    # Restart the same app with the same test PostgreSQL and photo directory.
    with TestClient(m.app) as restarted:
        login(restarted)
        assert restarted.get('/api/admin/reports/'+id+'/photos/'+path.name).content==backup
        assert len(PdfReader(io.BytesIO(restarted.get('/api/admin/reports/'+id+'/pdf').content)).pages)==1
    # Test snapshot copies and erasure reapplication exclusively in disposable DB.
    assert live.request('DELETE','/api/admin/reports/'+id,json={'confirm':'ELIMINAR '+id},headers=csrf).status_code==200
    ledger=live.get('/api/admin/export/erasures').json()
    with m.db() as c:
        c.execute('INSERT INTO reports(id,idempotency_key,request_hash,fingerprint,data,photos) VALUES (%s,%s,%s,%s,%s::jsonb,%s::jsonb)',(id,record['idempotency_key'],record['request_hash'],record['fingerprint'],json.dumps(record['data']),json.dumps(record['photos'])))
    path.write_bytes(backup)
    # Before reopening a restored snapshot, reapply the current erasure ledger.
    for entry in ledger:
        result=live.request('DELETE','/api/admin/reports/'+entry['id'],json={'confirm':'ELIMINAR '+entry['id']},headers=csrf)
        assert result.status_code==200
    assert not path.exists() and live.get('/api/admin/reports?archive=all').json()['total']==0

def test_migration_and_legacy_matching(live,data):
    legacy=m.validate(data)
    for field in ['family_name','total_people',*m.COUNT_FIELDS,'coordinates']:legacy.pop(field,None)
    with m.db() as c:
        c.execute((m.ROOT.parent/'migrations/004_household_lifecycle.sql').read_text())
        c.execute((m.ROOT.parent/'migrations/004_household_lifecycle.sql').read_text())
        c.execute('INSERT INTO reports(idempotency_key,request_hash,fingerprint,data) VALUES (%s,%s,%s,%s::jsonb)',(str(uuid.uuid4()),'legacy',m.signing(legacy['phone']+'|'+legacy['maps']),json.dumps(legacy)))
    with TestClient(m.app) as restarted:
        login(restarted);r=restarted.get('/api/admin/reports').json()['reports'][0]
        assert r['dedup_keys'] and r['data'].get('total_people') is None
        assert restarted.get('/api/admin/reports/'+r['id']+'/pdf').status_code==200
        assert submit(restarted,{**data,'family_name':'Nueva familia'}).json()['possible_duplicate']
