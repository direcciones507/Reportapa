import io, json, uuid, zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from xml.etree import ElementTree as ET
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import main as m
from app.report_statistics import statistics, recorded_needs, UNSPECIFIED
from app.report_excel import SHEETS
from test_app import live, data, submit

def record(identifier, family='Familia A', phone='61234567', status='pendiente', help_state='pendiente_confirmar', confirmed=False, needs=None, community='Comunidad A'):
    return dict(id=identifier, created_at=datetime(2026,10,11,12,tzinfo=timezone.utc), status=status,
        estado_ayuda=help_state, ayuda_actualizada_at=datetime(2026,10,12,tzinfo=timezone.utc) if confirmed else None,
        data=dict(name='Contacto privado', family_name=family, phone=phone, province='Veraguas', district='Santiago',
                  township='Santiago', community=community, damage_type='Paredes', needs=needs or []))

def test_distinct_units_states_and_explicit_needs():
    first=record('RPA-00000001',needs=['Agua potable','Pañales para bebés','Pañales para adultos'])
    second=record('RPA-00000002',status='visitado',needs=['Agua potable'])
    partial=record('RPA-00000003',family='Familia B',phone='69876543',help_state='parcial',confirmed=True,needs=['Colchones o lugar para dormir'])
    unconfirmed=record('RPA-00000004',family='',status='contactado',help_state='completa',needs=[])
    unconfirmed['data']['needs_detail']='Necesita alimentos y alojamiento';unconfirmed['data']['damage']='Vivienda destruida'
    rows=[first,second,partial,unconfirmed,first]
    before=json.dumps(rows,default=str,sort_keys=True)
    s=statistics(rows)
    assert s['reports']==4 and s['family_groups']==2 and s['unidentified_family_reports']==1
    assert s['housing_count'] is None and s['tracking']==dict(pending=1,followup=2,attended=1,unspecified=0)
    assert sum(s['tracking'].values())==s['reports']
    needs={r['category']:r for r in s['needs']}
    assert needs['Agua potable']['reports']==2 and needs['Agua potable']['family_groups']==1
    assert needs['Ropa y artículos básicos']['reports']==1 and needs['Ropa y artículos básicos']['family_groups']==1
    assert needs['Colchones o lugar para dormir']['reports']==1
    assert needs['Alojamiento temporal']['reports'] is None
    assert needs['Alimentos']['reports']==0 and needs[UNSPECIFIED]['reports']==1
    assert needs[UNSPECIFIED]['unidentified_family_reports']==1
    assert json.dumps(rows,default=str,sort_keys=True)==before

def test_hierarchy_legacy_and_unknown_categories():
    a=record('RPA-00000001',community='El Centro')
    b=record('RPA-00000002',community='El Centro');b['data']['district']='Otro distrito'
    b['data'].pop('family_name');b['data'].pop('damage_type');b['data']['needs']=['Categoría antigua',None]
    s=statistics([a,b]);assert len(s['communities'])==2
    assert sum(r['reports'] for r in s['communities'])==2
    assert recorded_needs(b)=={UNSPECIFIED}
    assert next(r for r in s['damages'] if r['type']==UNSPECIFIED)['reports']==1

def test_six_sheets_real_tables_chart_references_and_privacy():
    rows=[record('RPA-00000001',needs=['Agua potable']),record('RPA-00000002',family='Familia B',phone='69876543',status='visitado')]
    rows[0]['data'].update(name='=HYPERLINK("evil")',maps='https://maps.app.goo.gl/private',damage='PRIVATE_DAMAGE',needs_detail='PRIVATE_NEEDS')
    rows[0]['ayuda_observacion']='PRIVATE_NOTE'
    content=m.workbook(rows,json.dumps({'community':'Comunidad A'}),include_personal=False)
    wb=load_workbook(io.BytesIO(content));assert tuple(wb.sheetnames)==SHEETS
    assert all(sheet.tables for sheet in wb)
    assert wb['Resumen general']['B6'].value==2 and wb['Resumen general']['B7'].value==2
    assert wb['Detalle de reportes'].max_row==3
    assert isinstance(wb['Detalle de reportes']['B2'].value,datetime)
    text=str([list(s.values) for s in wb])
    for secret in ['HYPERLINK','61234567','PRIVATE_DAMAGE','PRIVATE_NEEDS','PRIVATE_NOTE','maps.app.goo.gl']:
        assert secret not in text
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        charts=[n for n in z.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')]
        tables=[n for n in z.namelist() if n.startswith('xl/tables/table') and n.endswith('.xml')]
        assert len(charts)==6 and len(tables)==9
        ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        for name in charts:
            xml=ET.fromstring(z.read(name));refs=xml.findall('.//c:f',ns)
            assert refs and all('!' in r.text and '[' not in r.text for r in refs)
            for r in refs:
                sheet,area=r.text.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
                assert sheet in wb.sheetnames and wb[sheet][area.replace('$','')]
            assert all(s.find('c:tx/c:v',ns) is not None for s in xml.findall('.//c:ser',ns))
        assert not any('externalLinks/' in n for n in z.namelist())
    private=load_workbook(io.BytesIO(m.workbook(rows,'todos',include_personal=True)))
    assert private['Detalle de reportes']['C2'].value=='=HYPERLINK("evil")'
    assert private['Detalle de reportes']['C2'].data_type=='s'
    anon=load_workbook(io.BytesIO(m.workbook(rows,json.dumps({'community':'Comunidad A'}),True,True)))
    anonymous=str([list(s.values) for s in anon])
    for secret in ['Comunidad A','Familia A','Familia B','61234567','HYPERLINK','PRIVATE_','maps.app.goo.gl']:
        assert secret not in anonymous

def test_empty_export_and_no_generated_needs():
    s=statistics([]);assert s['reports']==s['family_groups']==0 and not s['communities']
    wb=load_workbook(io.BytesIO(m.workbook([],'sin resultados',include_personal=False)))
    assert tuple(wb.sheetnames)==SHEETS and wb['Resumen general']['B6'].value==0
    row={r[0]:r for r in wb['Necesidades'].iter_rows(min_row=5,max_row=14,values_only=True)}
    assert row['Alojamiento temporal'][1]==UNSPECIFIED
    with zipfile.ZipFile(io.BytesIO(m.workbook([],'sin resultados',include_personal=False))) as z:
        assert not any(n.startswith('xl/charts/chart') for n in z.namelist())

def login(client):
    r=client.post('/api/admin/login',json={'username':'Ayudapanama','password':'test-password-123456'})
    assert r.status_code==200

def insert_fixture(row):
    with m.db() as c:
        c.execute('INSERT INTO reports(id,created_at,idempotency_key,request_hash,fingerprint,data,status,estado_ayuda,ayuda_actualizada_at) VALUES (%s,%s,%s,%s,%s,%s::jsonb,%s,%s,%s)',
            (row['id'],row['created_at'],str(uuid.uuid4()),str(uuid.uuid4()),str(uuid.uuid4()),json.dumps(row['data']),row['status'],row['estado_ayuda'],row['ayuda_actualizada_at']))

def test_filtered_panel_excel_same_snapshot_and_historical_readonly(live):
    rows=[record(f'RPA-{i:08d}',family='Familia '+str(i),phone=str(61230000+i)) for i in range(1,26)]
    rows[0]['created_at']=datetime(2026,10,11,4,59,59,tzinfo=timezone.utc)  # previous Panama date
    rows[1]['created_at']=datetime(2026,10,11,5,tzinfo=timezone.utc)
    rows[2]['created_at']=datetime(2026,10,12,4,59,59,tzinfo=timezone.utc)
    rows[3]['created_at']=datetime(2026,10,12,5,tzinfo=timezone.utc)  # next Panama date
    rows[4]['data']['province']='Coclé';rows[5]['data']['district']='Soná'
    rows[6]['data']['township']='Otro';rows[7]['data']['community']='Comunidad B'
    rows[8].update(status='visitado');rows[9].update(estado_ayuda='completa',ayuda_actualizada_at=datetime.now(timezone.utc))
    rows[10]['data']['needs']=['Agua potable'];rows[11]['data']['needs']=['Categoría antigua',None]
    for row in rows:insert_fixture(row)
    assert TestClient(m.app).get('/api/admin/reports?include_statistics=true').status_code==401
    login(live)
    original=m.db
    with original() as c:
        before=c.execute('SELECT * FROM reports ORDER BY id').fetchall()
    @contextmanager
    def readonly():
        with original() as c:
            c.execute('SET TRANSACTION READ ONLY');yield c
    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(m,'db',readonly)
        all_reports=live.get('/api/admin/reports?include_statistics=true').json()
        assert all_reports['total']==all_reports['statistics']['reports']==25
        assert len(all_reports['reports'])==20
        second=live.get('/api/admin/reports?include_statistics=true&page=2').json()
        assert len(second['reports'])==5 and second['statistics']['reports']==25
        checks=[({'province':'Coclé'},1),({'district':'Soná'},1),({'township':'Otro'},1),({'community':'Comunidad B'},1),
                ({'province':'Veraguas','district':'Santiago','township':'Santiago','community':'Comunidad A'},21),
                ({'tracking':'attended'},1),({'tracking':'followup'},1),({'status':'visitado'},1),
                ({'need':'Agua potable'},1),({'need':UNSPECIFIED},24),({'community':"' OR TRUE --"},0),
                ({'date_from':'2026-10-11','date_to':'2026-10-11'},23)]
        for filters,expected in checks:
            result=live.get('/api/admin/reports',params={**filters,'include_statistics':'true'})
            assert result.status_code==200,result.text
            d=result.json();assert d['total']==d['statistics']['reports']==expected
            excel=live.get('/api/admin/export/area',params=filters);assert excel.status_code==200
            assert excel.headers['content-type'].startswith('application/vnd.openxmlformats')
            wb=load_workbook(io.BytesIO(excel.content));assert wb['Resumen general']['B6'].value==expected
            detail=list(wb['Detalle de reportes'].iter_rows(min_row=2,values_only=True))
            assert len([r for r in detail if r[0]])==expected
            assert 'Teléfono' not in [c.value for c in wb['Detalle de reportes'][1]]
        full=load_workbook(io.BytesIO(live.get('/api/admin/export/area?include_personal=true').content))
        assert 'Teléfono' in [c.value for c in full['Detalle de reportes'][1]]
        for filters in [{'date_from':'2026-02-30'},{'date_from':'2026-10-12','date_to':'2026-10-11'},
                        {'tracking':'wrong'},{'status':'wrong'},{'date_to':'9999-12-31'}]:
            assert live.get('/api/admin/reports',params=filters).status_code==422
            assert live.get('/api/admin/export/area',params=filters).status_code==422
    with original() as c:assert c.execute('SELECT * FROM reports ORDER BY id').fetchall()==before

def test_citizen_report_without_maps_photo_or_needs_remains_valid(live,data):
    fields={**data,'maps':'','needs':[],'needs_detail':'','consent':'yes','idempotency_key':str(uuid.uuid4())}
    fields.pop('needs')
    r=live.post('/api/reports',data=fields,files={'photos':('',b'','application/octet-stream')})
    assert r.status_code==200 and r.json()['id'].startswith('RPA-')
    login(live);d=live.get('/api/admin/reports?include_statistics=true').json()
    assert d['statistics']['reports']==1
    assert next(n for n in d['statistics']['needs'] if n['category']==UNSPECIFIED)['reports']==1
