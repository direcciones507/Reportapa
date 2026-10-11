import os, json, uuid, hashlib, hmac, time, io, re, secrets, asyncio
from pathlib import Path
from urllib.parse import urlparse
from contextlib import asynccontextmanager
import psycopg
from psycopg.rows import dict_row
from fastapi import FastAPI, Request, HTTPException, Depends
from fastapi.responses import FileResponse, Response, JSONResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image, ImageOps, UnidentifiedImageError
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from .individual_pdf import individual_pdf
from .report_data import household, coordinates, matching_keys, COUNT_FIELDS
from .photo_maintenance import cleanup_photos

ROOT = Path(__file__).resolve().parent
NEEDS = ['Agua potable','Alimentos','Pañales para bebés','Pañales para adultos','Medicamentos','Ropa y cobijas','Colchones o lugar para dormir','Materiales para reparar la vivienda','Artículos de higiene','Otra necesidad']
FIELDS = ['name','phone','province','district','township','community','maps','damage_type','damage','needs_detail']
ADMIN_USERNAME = 'Ayudapanama'
HELP_STATES = {'pendiente_confirmar':'Pendiente de confirmar','sin_ayuda':'Sin ayuda recibida','parcial':'Ayuda parcial','completa':'Ayuda completa'}
NOTICE = 'Este es un registro ciudadano independiente. Reportar daños o necesidades no garantiza una inspección ni la entrega de ayuda. Si existe peligro inmediato, contacte a las autoridades de emergencia.'

def db():
    return psycopg.connect(os.environ['DATABASE_URL'], row_factory=dict_row)

def photo_dir():
    p = Path(os.getenv('PHOTO_DIR', '/data/private-photos')); p.mkdir(mode=0o700, parents=True, exist_ok=True); return p

def validate_runtime():
    if len(os.getenv('SESSION_SECRET','')) < 32 or len(os.getenv('ADMIN_PASSWORD','')) < 16:
        raise RuntimeError('Configure SESSION_SECRET (32+) y ADMIN_PASSWORD (16+)')
    if os.getenv('COOKIE_SECURE','true') not in ('true','false'):
        raise RuntimeError('COOKIE_SECURE debe ser true o false')
    if os.getenv('PUBLIC_REPORTS_ENABLED','false') not in ('true','false'):
        raise RuntimeError('PUBLIC_REPORTS_ENABLED debe ser true o false')
    if os.getenv('RAILWAY_ENVIRONMENT_ID'):
        mount=os.getenv('RAILWAY_VOLUME_MOUNT_PATH','')
        if not mount or not Path(mount).is_dir() or not Path(os.getenv('PHOTO_DIR','/data/private-photos')).resolve().is_relative_to(Path(mount).resolve()):
            raise RuntimeError('Configure PHOTO_DIR dentro del volumen privado persistente de Railway')
        if os.getenv('COOKIE_SECURE','true') != 'true':
            raise RuntimeError('Railway requiere COOKIE_SECURE=true y HTTPS')

@asynccontextmanager
async def lifespan(app):
    validate_runtime()
    with db() as c:
        c.execute("SELECT pg_advisory_xact_lock(hashtextextended('reportapa-migrations',2))")
        for migration in sorted((ROOT.parent/'migrations').glob('*.sql')):
            c.execute(migration.read_text())
    photo_dir()
    # Backfill matching keys from real legacy data; do not invent household counts.
    with db() as c:
        for report in c.execute("SELECT id,data FROM reports WHERE dedup_keys='{}'").fetchall():
            data=report['data'];data['coordinates']=coordinates(data['maps'])
            c.execute('UPDATE reports SET data=%s::jsonb,dedup_keys=%s WHERE id=%s',(json.dumps(data),matching_keys(data,signing),report['id']))
    await asyncio.to_thread(cleanup_photos,db,photo_dir())
    async def maintain():
        while True:
            await asyncio.sleep(30)
            try: await asyncio.to_thread(cleanup_photos,db,photo_dir())
            except (psycopg.Error,OSError): pass  # Retry after transient database interruption.
    worker=asyncio.create_task(maintain())
    try: yield
    finally:
        worker.cancel()
        try: await worker
        except asyncio.CancelledError: pass

app = FastAPI(lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)
app.mount('/static', StaticFiles(directory=ROOT/'static'), name='static')

@app.middleware('http')
async def headers(request, call_next):
    if request.method in ('POST','PATCH','DELETE'):
        maximum = 16*1024*1024 if request.url.path == '/api/reports' else 4096
        body = bytearray()
        async for chunk in request.stream():
            body.extend(chunk)
            if len(body) > maximum: return JSONResponse({'detail':'Solicitud demasiado grande'},status_code=413)
        request._body = bytes(body)
    response = await call_next(request)
    response.headers.update({'Cache-Control':'no-store','X-Content-Type-Options':'nosniff','X-Frame-Options':'DENY','X-Robots-Tag':'noindex, nofollow','Referrer-Policy':'no-referrer','Content-Security-Policy':"default-src 'self'; img-src 'self' blob:; style-src 'self'; script-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'"})
    return response

def signing(value):
    return hmac.new(os.environ['SESSION_SECRET'].encode(), value.encode(), hashlib.sha256).hexdigest()

def admin(request: Request):
    token = request.cookies.get('rpa_session','')
    try:
        expiry, nonce, signature = token.split('.')
        if int(expiry) < time.time() or not hmac.compare_digest(signing(expiry+'.'+nonce).encode(),signature.encode()): raise ValueError()
    except (ValueError, KeyError): raise HTTPException(401,'Sesión requerida')
    with db() as c:
        active=c.execute('SELECT username,credential_hash FROM admin_sessions WHERE token_hash=%s AND expires_at>now()', (hashlib.sha256(token.encode()).hexdigest(),)).fetchone()
    if not active or active['username'] != ADMIN_USERNAME or not hmac.compare_digest(active['credential_hash'],signing(os.environ['ADMIN_PASSWORD'])):
        raise HTTPException(401,'Sesión requerida')
    if request.method != 'GET' and not hmac.compare_digest(request.headers.get('x-csrf-token','').encode(), signing(token).encode()):
        raise HTTPException(403,'Solicitud inválida')
    return token

def limit(request, action, maximum):
    # Use the direct peer; do not trust client-supplied forwarding headers.
    key = signing(action+':'+(request.client.host if request.client else 'unknown'))
    with db() as c:
        r = c.execute("INSERT INTO request_limits(key,count) VALUES (%s,1) ON CONFLICT(key) DO UPDATE SET count=CASE WHEN request_limits.started_at < now()-interval '15 minutes' THEN 1 ELSE request_limits.count+1 END, started_at=CASE WHEN request_limits.started_at < now()-interval '15 minutes' THEN now() ELSE request_limits.started_at END RETURNING count",(key,)).fetchone()
        c.execute("DELETE FROM request_limits WHERE started_at < now()-interval '1 day'")
    if r['count'] > maximum: raise HTTPException(429,'Demasiados intentos. Espere 15 minutos.')

def validate(raw):
    data = {k: str(raw.get(k,'')).strip() for k in FIELDS}
    for k in FIELDS:
        if k not in ('needs_detail','damage_type') and not data[k]: raise HTTPException(422,'Complete los campos obligatorios')
        if len(data[k]) > (1500 if k in ('damage','needs_detail') else 300): raise HTTPException(422,'Campo demasiado largo')
    if not re.fullmatch(r'[+0-9() .-]+',data['phone']): raise HTTPException(422,'Teléfono de Panamá inválido')
    data['phone'] = re.sub(r'[^0-9]','',data['phone'])
    if len(data['phone']) == 11 and data['phone'].startswith('507'): data['phone'] = data['phone'][3:]
    if not re.fullmatch(r'[2-9][0-9]{6,7}',data['phone']): raise HTTPException(422,'Teléfono de Panamá inválido')
    try:
        u = urlparse(data['maps'])
        port = u.port
    except ValueError: raise HTTPException(422,'Enlace Maps inválido')
    if u.scheme != 'https' or u.username or u.password or port or u.hostname not in ('maps.app.goo.gl','goo.gl','www.google.com','google.com','maps.google.com'):
        raise HTTPException(422,'Use un enlace HTTPS compartido de Google Maps')
    if u.hostname in ('google.com','www.google.com') and not u.path.startswith('/maps'): raise HTTPException(422,'Enlace Maps inválido')
    if u.hostname == 'goo.gl' and not u.path.startswith('/maps/'): raise HTTPException(422,'Enlace Maps inválido')
    if not u.path.strip('/'): raise HTTPException(422,'Enlace Maps incompleto')
    needs = raw.get('needs',[])
    if not isinstance(needs,list) or any(n not in NEEDS for n in needs): raise HTTPException(422,'Necesidad inválida')
    data['needs'] = list(dict.fromkeys(needs))
    mode=str(raw.get('reporting_for_self','yes')).strip()
    if mode not in ('yes','no'): raise HTTPException(422,'Indique quién realiza el reporte')
    reporter=str(raw.get('reporter_name','')).strip()
    if len(reporter)>300: raise HTTPException(422,'Nombre de quien reporta demasiado largo')
    if mode=='no' and not reporter: raise HTTPException(422,'Indique el nombre de quien realiza el reporte')
    data['reporting_for_self']=mode
    data['reporter_name']=reporter if mode=='no' else ''
    data.update(household(raw));data['coordinates']=coordinates(data['maps'])
    return data

def clean_photo(content):
    if len(content) > 5*1024*1024: raise HTTPException(422,'Cada foto debe pesar menos de 5 MB')
    try:
        with Image.open(io.BytesIO(content)) as im:
            if im.format not in ('JPEG','PNG','WEBP') or im.width*im.height > 25_000_000: raise ValueError()
            im = ImageOps.exif_transpose(im); im.thumbnail((2000,2000))
            out = io.BytesIO(); im.convert('RGB').save(out,'JPEG',quality=85); return out.getvalue()
    except (UnidentifiedImageError,ValueError,OSError,Image.DecompressionBombError): raise HTTPException(422,'Foto inválida: use JPG, PNG o WebP')

@app.get('/')
def home(): return FileResponse(ROOT/'static/index.html')
@app.get('/admin')
def panel(): return FileResponse(ROOT/'static/admin.html')
@app.get('/health')
def health():
    with db() as c: c.execute('SELECT 1')
    return {'status':'ok'}
@app.get('/api/config')
def config(): return {'enabled':os.getenv('PUBLIC_REPORTS_ENABLED','false') == 'true','needs':NEEDS,'notice':NOTICE}
async def json_payload(request):
    try: return await request.json()
    except (ValueError,UnicodeDecodeError): raise HTTPException(422,'Solicitud inválida')

@app.post('/api/admin/login')
async def login(request: Request):
    limit(request,'login',10)
    data = await json_payload(request)
    if not isinstance(data,dict): raise HTTPException(422,'Solicitud inválida')
    username=data.get('username',''); password=data.get('password','')
    if not isinstance(username,str) or not isinstance(password,str): raise HTTPException(401,'Credenciales inválidas')
    user_ok=hmac.compare_digest(username.encode(),ADMIN_USERNAME.encode())
    password_ok=hmac.compare_digest(password.encode(),os.environ['ADMIN_PASSWORD'].encode())
    if not (user_ok and password_ok): raise HTTPException(401,'Credenciales inválidas')
    expiry=int(time.time()+3600)
    value = str(expiry)+'.'+secrets.token_hex(24); token = value+'.'+signing(value)
    with db() as c:
        c.execute('DELETE FROM admin_sessions WHERE expires_at<=now()')
        c.execute('INSERT INTO admin_sessions(token_hash,username,expires_at,credential_hash) VALUES (%s,%s,to_timestamp(%s),%s)',
                  (hashlib.sha256(token.encode()).hexdigest(),ADMIN_USERNAME,expiry,signing(os.environ['ADMIN_PASSWORD'])))
    response = JSONResponse({'csrf':signing(token)})
    response.set_cookie('rpa_session',token,httponly=True,secure=os.getenv('COOKIE_SECURE','true') == 'true',samesite='strict',max_age=3600,path='/'); return response
@app.get('/api/admin/session')
def session(token=Depends(admin)): return {'csrf':signing(token),'username':ADMIN_USERNAME}
@app.post('/api/admin/logout')
def logout(token=Depends(admin)):
    with db() as c: c.execute('DELETE FROM admin_sessions WHERE token_hash=%s',(hashlib.sha256(token.encode()).hexdigest(),))
    r = JSONResponse({'ok':True}); r.delete_cookie('rpa_session',path='/'); return r

@app.post('/api/reports')
async def submit(request: Request):
    if os.getenv('PUBLIC_REPORTS_ENABLED','false') != 'true': raise HTTPException(503,'Recepción de reportes aún no habilitada')
    limit(request,'submit',20)
    form = await request.form(max_files=3,max_fields=40)
    if form.get('website'): raise HTTPException(422,'Solicitud inválida')
    try: key = str(uuid.UUID(str(form.get('idempotency_key',''))))
    except ValueError: raise HTTPException(422,'Clave de envío inválida')
    raw = dict(form); raw['needs'] = form.getlist('needs'); data = validate(raw)
    if form.get('consent') != 'yes': raise HTTPException(422,'Debe autorizar el uso privado de los datos')
    photos = []
    for f in form.getlist('photos'):
        if hasattr(f,'read') and f.filename:
            photos.append(clean_photo(await f.read(5*1024*1024+1)))
    if len(photos)>3: raise HTTPException(422,'Máximo tres fotos')
    digest = hashlib.sha256(json.dumps(data,sort_keys=True).encode()+b''.join(hashlib.sha256(p).digest() for p in photos)).hexdigest()
    fingerprint = signing(data['phone']+'|'+data['maps'].rstrip('/'))
    dedup=matching_keys(data,signing)
    saved = []
    try:
        with db() as c:
            c.execute('SELECT pg_advisory_xact_lock(hashtextextended(%s,0))',(key,))
            for match_key in dedup:
                c.execute('SELECT pg_advisory_xact_lock(hashtextextended(%s,1))',(match_key,))
            existing = c.execute('SELECT id,request_hash FROM reports WHERE idempotency_key=%s',(key,)).fetchone()
            if existing:
                if existing['request_hash'] != digest: raise HTTPException(409,'La clave ya se utilizó con otro contenido')
                return {'id':existing['id'],'notice':NOTICE}
            identical=c.execute('SELECT 1 FROM reports WHERE request_hash=%s LIMIT 1',(digest,)).fetchone()
            if identical: raise HTTPException(409,'No se guardó otro reporte: el envío coincide exactamente con uno previo. Solicite revisión al administrador si necesita corregirlo.')
            duplicate = c.execute('SELECT 1 FROM reports WHERE dedup_keys && %s::text[] OR fingerprint=%s LIMIT 1',(dedup,fingerprint)).fetchone()
            if duplicate and form.get('new_damage') != 'yes':
                return JSONResponse({'possible_duplicate':True,'message':'Hay una coincidencia probable. Si son daños nuevos o corresponde a otra familia, confirme para enviar un nuevo reporte.'},status_code=409)
            for photo in photos:
                filename = uuid.uuid4().hex+'.jpg'; path = photo_dir()/filename
                descriptor=os.open(path,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
                saved.append(filename)
                with os.fdopen(descriptor,'wb') as output: output.write(photo)
            r = c.execute('INSERT INTO reports(idempotency_key,request_hash,fingerprint,data,photos,dedup_keys) VALUES (%s,%s,%s,%s::jsonb,%s::jsonb,%s) RETURNING id',(key,digest,fingerprint,json.dumps(data),json.dumps(saved),dedup)).fetchone()
        return {'id':r['id'],'notice':NOTICE}
    except Exception:
        pending=[]
        for filename in saved:
            try: (photo_dir()/filename).unlink(missing_ok=True)
            except OSError: pending.append(filename)
        if pending:
            try:
                with db() as c:
                    for filename in pending:
                        c.execute('INSERT INTO photo_cleanup(filename) VALUES (%s) ON CONFLICT DO NOTHING',(filename,))
            except psycopg.Error: pass  # Database outage: private orphan requires later reconciliation.
        raise

def rows(request, all_rows=False):
    clauses=[]; values=[]
    if not all_rows:
        archive=request.query_params.get('archive','active')
        if archive not in ('active','archived','all'): raise HTTPException(422,'Filtro de archivo inválido')
        if archive!='all': clauses.append('archived_at IS '+('NULL' if archive=='active' else 'NOT NULL'))
        for key in ('province','district','township','community','status','need','estado_ayuda'):
            value = request.query_params.get(key,'').strip()
            if not value: continue
            if key in ('status','estado_ayuda'):
                if key == 'estado_ayuda' and value not in HELP_STATES: raise HTTPException(422,'Estado de ayuda inválido')
                clauses.append(key+'=%s')
            elif key == 'need': clauses.append("data->'needs' ? %s")
            else: clauses.append('data->>%s = %s'); values.append(key)
            values.append(value)
    sql = 'SELECT * FROM reports'+(' WHERE '+' AND '.join(clauses) if clauses else '')+" ORDER BY data->>'province',data->>'district',data->>'township',created_at"
    with db() as c: return c.execute(sql,values).fetchall()
@app.get('/api/admin/reports')
def reports(request: Request, token=Depends(admin)):
    result=rows(request)
    try: page=max(1,int(request.query_params.get('page','1')))
    except ValueError: raise HTTPException(422,'Página inválida')
    return {'total':len(result),'reports':result[(page-1)*50:page*50]}
@app.patch('/api/admin/reports/{report_id}')
async def status(report_id: str, request: Request, token=Depends(admin)):
    payload = await json_payload(request)
    if not isinstance(payload,dict) or not payload or set(payload)-{'status','estado_ayuda','observacion','archived'}:
        raise HTTPException(422,'Cambio inválido')
    if 'status' in payload and payload['status'] not in ('pendiente','contactado','visitado'):
        raise HTTPException(422,'Estado operativo inválido')
    if 'estado_ayuda' in payload and (not isinstance(payload['estado_ayuda'],str) or payload['estado_ayuda'] not in HELP_STATES):
        raise HTTPException(422,'Estado de ayuda inválido')
    if 'observacion' in payload and 'estado_ayuda' not in payload: raise HTTPException(422,'Seleccione el estado de ayuda')
    if 'archived' in payload and not isinstance(payload['archived'],bool): raise HTTPException(422,'Archivo inválido')
    note = payload.get('observacion')
    if note is not None and (not isinstance(note,str) or len(note)>500): raise HTTPException(422,'Observación: máximo 500 caracteres')
    with db() as c:
        report=c.execute('SELECT * FROM reports WHERE id=%s FOR UPDATE',(report_id,)).fetchone()
        if not report: raise HTTPException(404)
        if 'archived' in payload:
            if payload['archived'] and not report['archived_at']:
                c.execute('UPDATE reports SET archived_at=now(),archived_by=%s WHERE id=%s',(ADMIN_USERNAME,report_id))
            elif not payload['archived']:
                c.execute('UPDATE reports SET archived_at=NULL,archived_by=NULL WHERE id=%s',(report_id,))
        if 'status' in payload: c.execute('UPDATE reports SET status=%s WHERE id=%s',(payload['status'],report_id))
        if 'estado_ayuda' in payload:
            note=report['ayuda_observacion'] if note is None else note.strip()
            if payload['estado_ayuda'] != report['estado_ayuda'] or note != report['ayuda_observacion']:
                # The authenticated session belongs exclusively to the configured owner.
                c.execute("""UPDATE reports SET estado_ayuda=%s, ayuda_observacion=%s,
                    ayuda_actualizada_at=now(), ayuda_admin=%s,
                    ayuda_historial=ayuda_historial || jsonb_build_array(jsonb_build_object(
                      'anterior',estado_ayuda,'estado',%s::text,'fecha',now(),
                      'observacion',%s::text,'administrador',%s::text)) WHERE id=%s""",
                    (payload['estado_ayuda'],note,ADMIN_USERNAME,payload['estado_ayuda'],note,ADMIN_USERNAME,report_id))
    return {'ok':True}

@app.delete('/api/admin/reports/{report_id}')
async def delete_report(report_id: str, request: Request, token=Depends(admin)):
    payload=await json_payload(request)
    if not isinstance(payload,dict) or payload.get('confirm') != 'ELIMINAR '+report_id:
        raise HTTPException(422,'Confirme escribiendo ELIMINAR seguido del número de reporte')
    with db() as c:
        report=c.execute('SELECT photos FROM reports WHERE id=%s FOR UPDATE',(report_id,)).fetchone()
        if not report: raise HTTPException(404)
        for name in report['photos']:
            c.execute('INSERT INTO photo_cleanup(filename) VALUES (%s) ON CONFLICT DO NOTHING',(name,))
        c.execute('INSERT INTO deleted_reports(id) VALUES (%s) ON CONFLICT DO NOTHING',(report_id,))
        c.execute('DELETE FROM reports WHERE id=%s',(report_id,))
    # After commit, the private routes can no longer serve the deleted photos.
    try:
        await asyncio.to_thread(cleanup_photos,db,photo_dir())
        with db() as c:
            pending=c.execute('SELECT count(*) AS n FROM photo_cleanup WHERE filename=ANY(%s)',(report['photos'],)).fetchone()['n']
    except (psycopg.Error,OSError):
        return JSONResponse({'deleted':True,'photos_pending':len(report['photos']),'message':'Datos eliminados. No se pudo verificar la limpieza física; se reintentará automáticamente.'},status_code=202)
    return JSONResponse({'deleted':True,'photos_pending':pending,'message':'Datos eliminados. '+('Limpieza física de fotos pendiente; se reintentará automáticamente.' if pending else 'Fotografías eliminadas.')},status_code=202 if pending else 200)

@app.get('/api/admin/reports/{report_id}/pdf')
def report_pdf(report_id: str, token=Depends(admin)):
    with db() as c: report=c.execute('SELECT * FROM reports WHERE id=%s',(report_id,)).fetchone()
    if not report: raise HTTPException(404)
    content=individual_pdf(report,photo_dir(),HELP_STATES)
    return Response(content,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="REPORTAPA_{report["id"]}_FICHA.pdf"'})
@app.get('/api/admin/reports/{report_id}/photos/{filename}')
def photo(report_id: str, filename: str, token=Depends(admin)):
    with db() as c: r=c.execute('SELECT photos FROM reports WHERE id=%s',(report_id,)).fetchone()
    if not r or filename not in r['photos'] or not re.fullmatch(r'[a-f0-9]{32}\.jpg',filename): raise HTTPException(404)
    path=photo_dir()/filename
    if not path.is_file(): raise HTTPException(404,'Foto no disponible')
    return FileResponse(path,media_type='image/jpeg')

def workbook(records, filters, anonymous=False):
    wb=Workbook(); sheet=wb.active; sheet.title='Reportes'
    columns=['ID de reporte','Fecha','Provincia','Distrito','Corregimiento','Comunidad','Daños','Necesidades','Detalle de necesidades','Estado'] if anonymous else ['ID de reporte','Fecha','Nombre del responsable / nombre original','Teléfono','Provincia','Distrito','Corregimiento','Comunidad','Google Maps','Tipo de daño','Descripción de daños','Necesidades','Detalle de necesidades','Estado']
    columns += ['Nombre de familia','Total de personas afectadas',*COUNT_FIELDS.values()]
    if not anonymous: columns += ['Reporte realizado por'] 
    if not anonymous: columns += ['Latitud','Longitud','Archivado']
    columns += ['Estado de ayuda','Fecha del último seguimiento']
    if not anonymous: columns.append('Observación de seguimiento')
    sheet.append(columns)
    for r in records:
        d=r['data']; base=[r['id'],r['created_at'].isoformat()]
        values = base+[d['province'],d['district'],d['township'],'','',', '.join(d['needs']),'',r['status']] if anonymous else base+[d['name'],d['phone'],d['province'],d['district'],d['township'],d['community'],d['maps'],d['damage_type'],d['damage'],', '.join(d['needs']),d['needs_detail'],r['status']]
        values += ['' if anonymous else d.get('family_name',''),d.get('total_people'),*[d.get(k) for k in COUNT_FIELDS]]
        if not anonymous: values += [d.get('reporter_name') or d.get('name','')]
        if not anonymous:
            point=d.get('coordinates') or {}
            values += [point.get('latitude'),point.get('longitude'),'Sí' if r.get('archived_at') else 'No']
        updated=r.get('ayuda_actualizada_at')
        values += [HELP_STATES[r.get('estado_ayuda','pendiente_confirmar')],updated.isoformat() if updated else '']
        if not anonymous: values.append(r.get('ayuda_observacion',''))
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=24
    meta=wb.create_sheet('Información'); meta.append(['Clasificación','Compartible anonimizado' if anonymous else 'CONFIDENCIAL — solo personal autorizado']); meta.append(['Fecha de corte',time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())]); meta.append(['Total de registros',len(records)]); meta.append(['Filtros',filters]); meta.append(['Aviso',NOTICE])
    for row in meta:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type='s'
    out=io.BytesIO(); wb.save(out); return out.getvalue()
@app.get('/api/admin/export/{kind}')
def export(kind: str, request: Request, token=Depends(admin)):
    if kind not in ('master','area','anonymous','summary','erasures'): raise HTTPException(404)
    if kind=='erasures':
        with db() as c: ledger=c.execute('SELECT id,deleted_at FROM deleted_reports ORDER BY deleted_at').fetchall()
        content=json.dumps([{'id':r['id'],'deleted_at':r['deleted_at'].isoformat()} for r in ledger])
        return Response(content,media_type='application/json',headers={'Content-Disposition':'attachment; filename=REPORTAPA_ELIMINACIONES.json'})
    records=rows(request,all_rows=kind=='master'); filters='Todos los reportes' if kind=='master' else json.dumps(dict(request.query_params),ensure_ascii=False)
    if kind=='summary':
        out=io.BytesIO(); pdf=canvas.Canvas(out); y=790
        for line in ['ReportaPA - resumen anonimizado',time.strftime('Fecha de corte: %Y-%m-%d'),f'Total de registros: {len(records)}']:
            pdf.drawString(40,y,line); y-=24
        counts={}
        for r in records:
            key=(r['data']['province'],r['status']); counts[key]=counts.get(key,0)+1
        for (province,state),count in sorted(counts.items()):
            if y<60: pdf.showPage(); y=790
            pdf.drawString(40,y,f'{province} / {state}: {count}'); y-=20
        pdf.save(); return Response(out.getvalue(),media_type='application/pdf',headers={'Content-Disposition':'attachment; filename="REPORTAPA_RESUMEN.pdf"'})
    scope='MASTER' if kind=='master' else re.sub(r'[^\w-]','_', '_'.join(request.query_params.get(k,'') for k in ('province','district','township'))).strip('_') or kind.upper()
    filename=f'REPORTAPA_{scope[:100]}_{time.strftime("%Y-%m-%d")}.xlsx'
    return Response(workbook(records,filters,kind=='anonymous'),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{filename}"'})
