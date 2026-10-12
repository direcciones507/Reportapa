"""Six-sheet extension of the existing openpyxl export, with native tables/charts."""
import io, json
from math import ceil
from collections import Counter
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from .report_data import COUNT_FIELDS
from .report_statistics import (statistics, unique_records, recorded_needs, tracking_group,
                                TRACKING_LABELS, LOCATION_FIELDS, UNSPECIFIED)

PANAMA = timezone(timedelta(hours=-5))
SHEETS = ('Resumen general', 'Comunidades', 'Tipos de daños', 'Necesidades', 'Seguimiento', 'Detalle de reportes')

def local_date(value):
    return value.astimezone(PANAMA).replace(tzinfo=None) if isinstance(value, datetime) else None

def safe_cell(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'  # Untrusted source text is never an executable Excel formula.
    if isinstance(value, datetime):
        cell.number_format = 'dd/mm/yyyy hh:mm'
    elif isinstance(value, int):
        cell.number_format = '#,##0'

def native_table(sheet, name, headers, data, row=4):
    for index, values in enumerate([headers, *data], row):
        for col, value in enumerate(values, 1):
            cell = sheet.cell(index, col); safe_cell(cell, value)
            cell.font = Font(name='Aptos', size=11, color='173044')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        if index > row:
            lines=max((ceil(len(str(v or ''))/24) for v in values),default=1)
            sheet.row_dimensions[index].height=min(110,max(24,lines*15+6))
    # A real table remains available even when filters match no reports. The blank row is not a report.
    last = row + max(1, len(data))
    table = Table(displayName=name, ref=f'A{row}:{get_column_letter(len(headers))}{last}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    for cell in sheet[row]:
        if cell.column <= len(headers):
            cell.font = Font(name='Aptos', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='123F58')
    sheet.row_dimensions[row].height = 50 if max(map(len,headers)) > 25 else 34
    return row + 1, row + len(data)

def chart(sheet, title, first, last, columns, category, anchor, unit='Reportes'):
    if last < first:
        sheet[anchor] = 'Sin reportes para los filtros seleccionados.'
        return
    values=[sheet.cell(row,col).value for row in range(first,last+1) for col in columns]
    maximum=max((v for v in values if isinstance(v,(int,float))),default=0)
    if not maximum:
        sheet[anchor] = 'Sin cantidades positivas para graficar con estos filtros.'
        return
    c = BarChart(); c.type = 'bar'; c.style = 10; c.title = title
    # openpyxl names the axes by class, even when a horizontal bar swaps their display positions.
    c.x_axis.title = 'Categoría'; c.y_axis.title = unit
    c.y_axis.numFmt = '0'; c.y_axis.majorUnit = max(1,ceil(maximum/5))
    c.y_axis.scaling.min = 0; c.y_axis.scaling.max = maximum+c.y_axis.majorUnit
    c.width = 21; c.height = 10; c.grouping = 'clustered'; c.overlap = 0
    for col in columns:
        c.add_data(Reference(sheet, min_col=col, min_row=first-1, max_row=last), titles_from_data=True)
        c.series[-1].tx = SeriesLabel(v=str(sheet.cell(first-1,col).value))
    c.set_categories(Reference(sheet, min_col=category, min_row=first, max_row=last))
    if len(columns) == 1: c.legend = None
    sheet.add_chart(c, anchor)

def statistical_workbook(records, filters, help_states, anonymous=False, include_personal=False):
    records = unique_records(records); stats = statistics(records, anonymous)
    wb = Workbook(); wb.remove(wb.active)
    for name in SHEETS:
        s = wb.create_sheet(name); s.sheet_view.showGridLines = False; s.freeze_panes = 'A5'
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.orientation = 'landscape'; s.page_setup.paperSize = s.PAPERSIZE_A4
        s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
        s.print_options.horizontalCentered = True
        s.oddFooter.center.text = 'ReportaPA · &P / &N'
        for col in range(1, 25): s.column_dimensions[get_column_letter(col)].width = 22
        s['A1'] = 'ReportaPA · ' + name; s['A1'].font = Font(name='Aptos', size=18, bold=True, color='123F58')
        s.row_dimensions[1].height = 30
    summary = wb[SHEETS[0]]
    summary.merge_cells('A2:D2'); summary['A2'] = 'Información agregada anonimizada' if anonymous else 'Informe administrativo privado'
    summary['A2'].alignment = Alignment(wrap_text=True); summary.row_dimensions[2].height = 30
    summary['A3'] = 'Fecha de corte (Panamá)'; safe_cell(summary['B3'], local_date(datetime.now(timezone.utc)))
    indicators = [['Total de reportes', stats['reports'], 'Reportes'],
        ['Grupos familiares declarados', stats['family_groups'], 'Grupos, no familias únicas verificadas'],
        ['Reportes sin identificación familiar', stats['unidentified_family_reports'], 'Reportes'],
        ['Viviendas únicas', 'No determinadas', 'No existe identificador de vivienda'],
        ['Comunidades', len(stats['communities']), 'Sectores con reportes']]
    native_table(summary, 'IndicadoresReportaPA', ['Indicador', 'Cantidad', 'Unidad'], indicators, row=5)
    state_rows = [[TRACKING_LABELS[k], v] for k, v in stats['tracking'].items()]
    first, last = native_table(summary, 'EstadosReportaPA', ['Seguimiento', 'Reportes'], state_rows, row=13)
    chart(summary, 'Seguimiento de reportes', first, last, [2], 1, 'F4')
    scopes = []
    for size, level in ((1, 'Provincia'), (2, 'Distrito'), (3, 'Corregimiento')):
        counts = Counter()
        for r in stats['communities']:
            loc = tuple(r['location'].values()); counts[loc[:size]] += r['reports']
        for loc, n in sorted(counts.items()): scopes.append([level, ' / '.join(loc), n])
    first, last = native_table(summary, 'UbicacionesReportaPA', ['Nivel', 'Ubicación', 'Reportes'], scopes, row=21)
    province_last = first + sum(r[0] == 'Provincia' for r in scopes) - 1
    chart(summary, 'Reportes por provincia', first, province_last, [3], 2, 'F25')
    note_row = max(25, last+3)
    if anonymous:
        try:
            selection = json.loads(filters)
            selection.pop('community', None); selection.pop('include_personal', None)
            filters = json.dumps(selection, ensure_ascii=False)
        except (ValueError, TypeError, AttributeError):
            filters = 'Filtros del informe; comunidad omitida'
    for label, value in [('Filtros aplicados', filters), ('Conteo de familias', stats['family_method']),
                         ('Atención registrada', stats['attention_method']), ('Necesidades', stats['needs_method'])]:
        safe_cell(summary.cell(note_row, 1), label)
        summary.cell(note_row, 1).alignment = Alignment(vertical='top',wrap_text=True)
        summary.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=4)
        safe_cell(summary.cell(note_row, 2), value); summary.cell(note_row, 2).alignment = Alignment(wrap_text=True, vertical='top')
        summary.row_dimensions[note_row].height = 68; note_row += 1

    community_rows = []
    followup_rows = []
    for r in stats['communities']:
        loc = list(r['location'].values()); label = ' / '.join(loc)
        community_rows.append([*loc, label, r['reports'], r['family_groups'], r['pending'], r['followup'],
            r['attended'], r['partial_attention'], r['complete_attention'], r['unspecified'],
            r['unidentified_family_reports'], ', '.join(f'{k}: {v}' for k, v in sorted(r['damages'].items())),
            ', '.join(f'{k}: {v}' for k, v in sorted(r['needs'].items()))])
        followup_rows.append([label, r['pending'], r['followup'], r['attended'], r['partial_attention'],
                              r['complete_attention'], r['unspecified'], r['reports']])
    communities = wb[SHEETS[1]]
    communities['A2'] = 'Cada fila corresponde a una comunidad dentro de su provincia, distrito y corregimiento.'
    first, last = native_table(communities, 'ComunidadesReportaPA', ['Provincia', 'Distrito', 'Corregimiento',
        'Comunidad', 'Ubicación completa', 'Reportes', 'Grupos familiares', 'Pendientes', 'En seguimiento',
        'Con atención registrada', 'Atención parcial', 'Atención completa', 'Estado no especificado',
        'Sin identificación familiar', 'Tipos de daños (reportes)', 'Necesidades (reportes)'], community_rows)
    chart(communities, 'Comunidades con más reportes (máximo 10)', first, min(last, first+9), [6], 5, f'A{max(8,last+4)}')
    damages = wb[SHEETS[2]]; damages['A2'] = 'Tipo seleccionado; no se interpreta la descripción libre de daños.'
    first, last = native_table(damages, 'DanosReportaPA', ['Tipo de daño', 'Reportes'],
                              [[r['type'], r['reports']] for r in stats['damages']])
    chart(damages, 'Tipos de daños registrados', first, last, [2], 1, 'E4')

    needs = wb[SHEETS[3]]
    needs['A2'] = 'Categorías múltiples. No sumar columnas como familias únicas. Sin datos suficientes: No especificada.'
    # The unavailable category is visibly preserved, but never plotted as a fabricated zero.
    need_rows = [[r['category'], r['family_groups'] if r['family_groups'] is not None else UNSPECIFIED,
                  r['reports'] if r['reports'] is not None else UNSPECIFIED,
                  r['unidentified_family_reports'] if r['unidentified_family_reports'] is not None else UNSPECIFIED,
                  r['availability']] for r in stats['needs'] if r['reports'] is not None]
    need_rows += [['Alojamiento temporal', UNSPECIFIED, UNSPECIFIED, UNSPECIFIED, 'Sin categoría inequívoca registrada']]
    first, last = native_table(needs, 'NecesidadesReportaPA', ['Necesidad', 'Grupos familiares', 'Reportes',
        'Reportes sin identificación familiar', 'Información disponible'], need_rows)
    chart(needs, 'Necesidades: grupos familiares y reportes', first, last-1, [2,3], 1, 'G4', 'Cantidad')
    detail_row = max(28, last+4)
    native_table(needs, 'NecesidadesComunidadesReportaPA', ['Provincia', 'Distrito', 'Corregimiento', 'Comunidad',
        'Necesidad', 'Grupos familiares', 'Reportes', 'Sin identificación familiar'],
        [[*r['location'].values(), r['category'], r['family_groups'], r['reports'], r['unidentified_family_reports']]
         for r in stats['needs_by_community']], row=detail_row)
    followup = wb[SHEETS[4]]; followup['A2'] = 'Atención parcial/completa con fecha confirmada. Contacto o evaluación no confirma ayuda.'
    first, last = native_table(followup, 'SeguimientoReportaPA', ['Comunidad y ubicación', 'Pendientes',
        'En seguimiento', 'Con atención registrada', 'Atención parcial', 'Atención completa',
        'Estado no especificado', 'Total de reportes'], followup_rows)
    chart(followup, 'Seguimiento por comunidad (máximo 10)', first, min(last, first+9), [2,3,4], 1, f'A{max(8,last+4)}')

    detail = wb[SHEETS[5]]; detail.delete_rows(1, 3); detail.freeze_panes = 'C2'
    personal = include_personal and not anonymous
    if personal:
        columns = ['ID de reporte', 'Fecha', 'Nombre del responsable / nombre original', 'Teléfono',
            'Provincia', 'Distrito', 'Corregimiento', 'Comunidad', 'Google Maps', 'Tipo de daño',
            'Descripción de daños', 'Necesidades', 'Detalle de necesidades', 'Estado',
            'Nombre de familia', 'Total de personas afectadas', *COUNT_FIELDS.values(),
            'Reporte realizado por', 'Latitud', 'Longitud', 'Archivado', 'Estado de ayuda',
            'Fecha del último seguimiento', 'Observación de seguimiento']
    else:
        columns = ['ID de reporte', 'Fecha', 'Provincia', 'Distrito', 'Corregimiento', 'Comunidad',
                   'Tipo de daño', 'Estado', 'Necesidades', 'Seguimiento', 'Estado de ayuda',
                   'Fecha del último seguimiento', 'Archivado']
    detail_rows = []
    safe_damage_types = ('Paredes', 'Techo', 'Piso', 'Varios daños', 'Otro')
    for r in records:
        d = r['data']; updated = local_date(r.get('ayuda_actualizada_at'))
        raw_needs=d.get('needs')
        supplied_needs=', '.join(n for n in raw_needs if isinstance(n,str)) if isinstance(raw_needs,list) else UNSPECIFIED
        help_label = help_states.get(r.get('estado_ayuda', 'pendiente_confirmar'), UNSPECIFIED)
        if r.get('estado_ayuda') in ('parcial','completa') and not updated:
            help_label += ' (sin confirmación fechada)'
        if personal:
            point = d.get('coordinates') or {}
            values = [r['id'], local_date(r['created_at']), d.get('name', ''), d.get('phone', ''),
                *(d.get(k, UNSPECIFIED) for k in LOCATION_FIELDS), d.get('maps', ''), d.get('damage_type', ''),
                d.get('damage', ''), supplied_needs, d.get('needs_detail', ''), r['status'],
                d.get('family_name', ''), d.get('total_people'), *(d.get(k) for k in COUNT_FIELDS),
                d.get('reporter_name') or d.get('name', ''), point.get('latitude'), point.get('longitude'),
                'Sí' if r.get('archived_at') else 'No', help_label, updated, r.get('ayuda_observacion', '')]
        else:
            loc = [d.get(k) or UNSPECIFIED for k in LOCATION_FIELDS]
            damage = d.get('damage_type') or UNSPECIFIED
            if anonymous:
                loc[3] = 'Comunidad omitida'
                if damage not in safe_damage_types: damage = UNSPECIFIED
            values = [r['id'], local_date(r['created_at']), *loc, damage, r.get('status') or UNSPECIFIED,
                      ', '.join(sorted(recorded_needs(r))), TRACKING_LABELS[tracking_group(r)],
                      help_label, updated, 'Sí' if r.get('archived_at') else 'No']
        detail_rows.append(values)
    native_table(detail, 'DetalleReportaPA', columns, detail_rows, row=1)
    for sheet in wb:
        sheet.print_title_rows = '1:4' if sheet != detail else '1:1'
        sheet.column_dimensions['A'].width = 32
        if sheet not in (summary, detail):
            sheet['A2'].alignment = Alignment(wrap_text=True)
            sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
            sheet.row_dimensions[2].height = 45
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
    out = io.BytesIO(); wb.save(out); return out.getvalue()
